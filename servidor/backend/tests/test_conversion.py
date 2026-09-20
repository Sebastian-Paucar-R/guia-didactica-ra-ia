import zipfile

import pytest

from app.services.conversion_service import ConversionError, convertir


def _docx(ruta, parrafos):
    cuerpo = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in parrafos)
    with zipfile.ZipFile(ruta, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                   "</Types>")
        z.writestr("_rels/.rels",
                   '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                   "</Relationships>")
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   f"<w:body>{cuerpo}</w:body></w:document>")


def _pptx(ruta, titulo, cuerpo):
    from pptx import Presentation

    prs = Presentation()
    diapositiva = prs.slides.add_slide(prs.slide_layouts[1])
    diapositiva.shapes.title.text = titulo
    diapositiva.placeholders[1].text = cuerpo
    prs.save(ruta)


def _convertir(origen, tmp_path):
    return convertir(origen, tmp_path / "pdf", tmp_path / "markdown")


def test_md_se_copia_igual_y_genera_pdf(tmp_path):
    origen = tmp_path / "norma.md"
    origen.write_text("# Norma ISO\n\nGestión de la **calidad**:\n\n- uno\n- dos\n\n| a | b |\n|---|---|\n| 1 | 2 |\n",
                      encoding="utf-8")
    r = _convertir(origen, tmp_path)
    assert r.markdown.read_bytes() == origen.read_bytes()
    assert r.pdf.read_bytes().startswith(b"%PDF")


def test_txt_a_md_y_pdf(tmp_path):
    origen = tmp_path / "notas.txt"
    origen.write_text("Línea con ñ y acentos: información.", encoding="cp1252")
    r = _convertir(origen, tmp_path)
    assert "información" in r.markdown.read_text(encoding="utf-8")
    assert r.pdf.read_bytes().startswith(b"%PDF")


def test_pdf_se_copia_igual_y_se_extrae_markdown(tmp_path):
    md = tmp_path / "origen.md"
    md.write_text("# Sistema de Gestión\n\nLa información es un activo.\n", encoding="utf-8")
    pdf = _convertir(md, tmp_path).pdf
    r = convertir(pdf, tmp_path / "pdf2", tmp_path / "markdown2")
    assert r.pdf.read_bytes() == pdf.read_bytes()
    texto = r.markdown.read_text(encoding="utf-8")
    assert "Gestión" in texto and "información" in texto


def test_docx(tmp_path):
    origen = tmp_path / "informe.docx"
    _docx(origen, ["Auditoría interna", "Mejora continua del proceso"])
    r = _convertir(origen, tmp_path)
    texto = r.markdown.read_text(encoding="utf-8")
    assert "Auditoría interna" in texto and "Mejora continua" in texto
    assert r.pdf.read_bytes().startswith(b"%PDF")


def test_pptx(tmp_path):
    origen = tmp_path / "clase.pptx"
    _pptx(origen, "Gestión de riesgos", "Identificar y tratar riesgos")
    r = _convertir(origen, tmp_path)
    assert "Gestión de riesgos" in r.markdown.read_text(encoding="utf-8")
    assert r.pdf.read_bytes().startswith(b"%PDF")


@pytest.mark.parametrize("nombre,contenido,mensaje", [
    ("virus.exe", b"MZ", "no soportada"),
    ("vacio.txt", b"   \n", "sin texto"),
    ("roto.docx", b"esto no es un docx", "DOCX válido"),
    ("roto.pdf", b"esto no es un pdf", "PDF válido"),
    ("roto.pptx", b"esto no es un pptx", "PPTX válido"),
    ("binario.txt", b"abc\x00\x01\x02", "binarios"),
])
def test_errores_controlados_y_sin_parciales(tmp_path, nombre, contenido, mensaje):
    origen = tmp_path / nombre
    origen.write_bytes(contenido)
    with pytest.raises(ConversionError, match=mensaje):
        _convertir(origen, tmp_path)
    for carpeta in ("pdf", "markdown"):
        assert not list((tmp_path / carpeta).glob("*")), "no deben quedar archivos parciales"
